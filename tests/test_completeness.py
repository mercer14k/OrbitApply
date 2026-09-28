from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock

import fitz
from docx import Document
from test_core import fixture_job, fixture_profile

from jobtailor.ai_tasks import (
    audit_and_repair_draft,
    extract_candidate_profile,
    tailor_resume,
)
from jobtailor.ats_resume import (
    experience_bullets,
    render_ats_pdf,
    resume_blocks,
    verify_pdf,
)
from jobtailor.models import (
    CandidateProfile,
    ExperienceEdit,
    TailoredBullet,
    TailoredResume,
)
from jobtailor.ollama_client import OllamaClient, OllamaError
from jobtailor.requirements_analysis import Requirement, RequirementAnalysis
from jobtailor.resume_parser import (
    ResumeParseError,
    bounded_chunks,
    parse_resume_bytes,
    profile_coverage,
    structured_profile,
)


def long_resume() -> str:
    # Synthetic, format-equivalent fixture. Never bundle a real person's resume.
    text = "Taylor Morgan\nChicago, IL | 555-0100 | taylor@example.com | LinkedIn: Taylor Morgan\nSUMMARY:\nSupply chain professional focused on planning and reporting.\nPROFESSIONAL EXPERIENCE:\n"
    for index, count in enumerate([10, 7, 8, 6, 4]):
        text += f"Employer {index}, IL, USA     January {2020 - index} - December {2021 - index}\nAnalyst {index}\n"
        for bullet in range(count):
            text += f"• Built reporting workflow {index}{bullet} for inventory\n  and procurement teams.\n"
    text += "EDUCATION:\n"
    for index in range(3):
        text += f"University {index}    January {2010 + index} - May {2012 + index}\nMaster of Science, Engineering\n"
    text += "SKILLS:\n● Tools: SQL, Power BI, Excel\n● Certified Quality Practitioner\n"
    return text


class CompletenessTests(unittest.TestCase):
    def test_all_roles_bullets_and_degrees_without_model(self):
        client = Mock()
        text = long_resume()
        profile = extract_candidate_profile(client, text)
        client.chat_json.assert_not_called()
        self.assertEqual(
            [len(e.bullets) for e in profile.experiences], [10, 7, 8, 6, 4]
        )
        self.assertEqual(len(profile.education), 3)
        self.assertEqual(profile.education[0].start_date, "January 2010")
        self.assertEqual(profile.skills, ["SQL", "Power BI", "Excel"])
        self.assertEqual(len(profile.certifications), 1)
        self.assertEqual(profile_coverage(text, profile).missing_lines, [])

    def test_coverage_detects_missing_last_page(self):
        text = long_resume()
        profile = structured_profile(text)
        profile.experiences = profile.experiences[:2]
        profile.education = []
        report = profile_coverage(text, profile)
        self.assertEqual(report.source_bullets, 35)
        self.assertEqual(report.profile_bullets, 17)
        self.assertTrue(report.missing_lines)
        self.assertTrue(report.warnings)

    def test_chunks_do_not_truncate_long_input(self):
        text = "A" * 30000 + "\nLAST EDUCATION ENTRY\n"
        chunks = bounded_chunks(text)
        self.assertEqual("".join(chunks), text)
        self.assertTrue(all(len(c) <= 2400 for c in chunks))

    def test_fallback_reads_last_fragment(self):
        client = Mock()
        client.chat_json.return_value = CandidateProfile()
        text = "Unusual layout\n" + "A" * 31000 + "\nLAST EDUCATION ENTRY"
        extract_candidate_profile(client, text)
        prompts = "".join(call.args[2] for call in client.chat_json.call_args_list)
        self.assertIn("LAST EDUCATION ENTRY", prompts)
        self.assertGreater(client.chat_json.call_count, 10)

    def test_mixed_scan_pdf_refused_not_partially_parsed(self):
        with fitz.open() as document:
            document.new_page().insert_text((40, 40), "Readable first page " * 8)
            document.new_page()  # Represents an image-only/unreadable page.
            with self.assertRaisesRegex(ResumeParseError, "Page 2"):
                parse_resume_bytes("mixed.pdf", document.tobytes())

    def test_two_page_text_pdf_retains_last_page(self):
        with fitz.open() as document:
            document.new_page().insert_text(
                (40, 40), "First employer and duties are here. " * 4
            )
            document.new_page().insert_text(
                (40, 40), "Final employer education and skills are here. " * 3
            )
            text = parse_resume_bytes("two.pdf", document.tobytes())
            self.assertIn("Final employer education", text)

    def test_pdf_embedded_link_retained(self):
        with fitz.open() as document:
            page = document.new_page()
            page.insert_text(
                (40, 40),
                "Taylor Morgan | LinkedIn: Taylor Morgan | Contact details and professional profile",
            )
            page.insert_link(
                {
                    "kind": fitz.LINK_URI,
                    "from": fitz.Rect(40, 25, 150, 45),
                    "uri": "https://www.linkedin.com/in/taylor/",
                }
            )
            text = parse_resume_bytes("linked.pdf", document.tobytes())
            self.assertIn("https://www.linkedin.com/in/taylor/", text)

    def test_encrypted_pdf_refused(self):
        with fitz.open() as document:
            document.new_page().insert_text(
                (40, 40), "Protected source with several fields and accomplishments."
            )
            data = document.tobytes(
                encryption=fitz.PDF_ENCRYPT_AES_256,
                owner_pw="test-owner",
                user_pw="test-user",
            )
            with self.assertRaisesRegex(ResumeParseError, "password protected"):
                parse_resume_bytes("locked.pdf", data)

    def test_docx_header_and_table_order(self):
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = "Contact: Taylor Morgan"
        doc.add_paragraph("Before the table, summary and experience overview.")
        doc.add_table(rows=1, cols=1).cell(
            0, 0
        ).text = "Inside the table, detailed employment evidence."
        doc.add_paragraph("After the table, education and other qualifications.")
        doc.add_paragraph(
            "Inventory reporting and process improvement", style="List Bullet"
        )
        buffer = io.BytesIO()
        doc.save(buffer)
        text = parse_resume_bytes("resume.docx", buffer.getvalue())
        self.assertIn("Contact: Taylor Morgan", text)
        self.assertLess(text.index("Inside the table"), text.index("After the table"))
        self.assertIn("• Inventory", text)

    def test_partial_tailoring_keeps_all_bullets(self):
        profile = fixture_profile()
        draft = TailoredResume(
            experience_edits=[
                ExperienceEdit(
                    experience_index=0,
                    bullets=[
                        TailoredBullet(
                            text="Reworded dashboard bullet.", source_bullet_indices=[0]
                        )
                    ],
                )
            ]
        )
        output = experience_bullets(profile, draft, 0)
        self.assertEqual(len(output), 3)
        self.assertEqual(output[1:], profile.experiences[0].bullets[1:])

    def test_combined_bullets_do_not_replace_source(self):
        profile = fixture_profile()
        draft = TailoredResume(
            experience_edits=[
                ExperienceEdit(
                    experience_index=0,
                    bullets=[
                        TailoredBullet(
                            text="Vague combined achievement",
                            source_bullet_indices=[0, 1],
                        )
                    ],
                )
            ]
        )
        self.assertEqual(
            experience_bullets(profile, draft, 0), profile.experiences[0].bullets
        )

    def test_all_skills_survive_highlighting(self):
        profile = fixture_profile()
        blocks = resume_blocks(profile, TailoredResume(highlighted_skills=["SQL"]))
        text = "\n".join(t for _, t in blocks)
        for skill in profile.skills:
            self.assertIn(skill, text)

    def test_length_stop_rejects_even_valid_partial_json(self):
        client = OllamaClient()
        client._request = Mock(
            return_value={"done_reason": "length", "message": {"content": "{}"}}
        )
        with self.assertRaisesRegex(OllamaError, "output limit"):
            client.chat_json(CandidateProfile, "Extract", "Short source")

    def test_large_prompt_fails_before_request(self):
        client = OllamaClient(context_window=4096)
        client._request = Mock()
        with self.assertRaisesRegex(OllamaError, "too large"):
            client.chat_json(CandidateProfile, "Extract", "A" * 30000)
        client._request.assert_not_called()

    def test_ats_round_trip_all_roles_and_numbers(self):
        profile = structured_profile(long_resume())
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "resume.pdf"
            report = render_ats_pdf(profile, TailoredResume(), path)
            self.assertTrue(report["passed"])
            with fitz.open(path) as document:
                text = "\n".join(page.get_text() for page in document)
                self.assertTrue(
                    any(
                        "Vera" in str(font)
                        for page in document
                        for font in page.get_fonts()
                    )
                )
            self.assertIn("Employer 4", text)
            self.assertIn("University 2", text)
            with self.assertRaisesRegex(ValueError, "missing or out of order"):
                verify_pdf(path, [("body", "This text was never in the resume")])

    def test_tailoring_and_audit_use_small_batches(self):
        profile = structured_profile(long_resume())
        client = Mock()

        from jobtailor.evidence_tailoring import EvidenceDecision

        def reply(model, system, prompt, **kwargs):
            self.assertLess(len(prompt), 4000)
            self.assertIs(model, EvidenceDecision)
            return EvidenceDecision(
                support="not_established",
                source_quote="",
                reason="No supported match in fixture",
            )

        client.chat_json.side_effect = reply
        analysis = RequirementAnalysis(
            requirements=[
                Requirement(kind="responsibility", quote="build Power BI dashboards")
            ]
        )
        tailored = tailor_resume(
            client, profile, fixture_job(), [], requirements=analysis
        )
        audited, report = audit_and_repair_draft(client, profile, tailored)
        self.assertTrue(report.passed)
        self.assertTrue(audited.tailoring_report["status"].startswith("Needs review"))
        self.assertFalse(audited.experience_edits)
        for i in range(5):
            self.assertEqual(
                experience_bullets(profile, audited, i), profile.experiences[i].bullets
            )


if __name__ == "__main__":
    unittest.main()
