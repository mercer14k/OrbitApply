from __future__ import annotations

import csv
import html
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from .ats_resume import experience_bullets, render_ats_pdf, resume_blocks
from .models import (
    CandidateProfile,
    DraftValidation,
    ExportResult,
    JobListing,
    TailoredResume,
)
from .requirements_analysis import RequirementAnalysis, tailoring_changes
from .utils import ensure_unique_folder, escape_paragraph, slug_component, write_json

ACCENT = RGBColor(31, 78, 121)


def _configure_docx(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(2)
    styles["Normal"].paragraph_format.line_spacing = 1.0
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(18)
    styles["Title"].font.bold = True
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(10.5)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].font.color.rgb = ACCENT
    styles["Heading 1"].paragraph_format.space_before = Pt(5)
    styles["Heading 1"].paragraph_format.space_after = Pt(2)


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _section_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text.upper(), style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True
    bottom_border = OxmlElement("w:pBdr")
    border = OxmlElement("w:bottom")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "5")
    border.set(qn("w:space"), "1")
    border.set(qn("w:color"), "9EADBA")
    bottom_border.append(border)
    paragraph._p.get_or_add_pPr().append(bottom_border)


def _experience_bullets(
    profile: CandidateProfile, tailored: TailoredResume, experience_index: int
) -> list[str]:
    return experience_bullets(profile, tailored, experience_index)


def render_resume_docx(
    profile: CandidateProfile,
    tailored: TailoredResume,
    destination: Path,
) -> None:
    document = Document()
    _configure_docx(document)
    name = profile.contact.name or "Candidate"
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(name)

    contact_items = [
        profile.contact.location,
        profile.contact.phone,
        profile.contact.email,
        profile.contact.linkedin,
        profile.contact.website,
    ]
    contact = document.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(3)
    contact.add_run(
        "  |  ".join(item for item in contact_items if item)
    ).font.size = Pt(8.5)

    if tailored.target_title:
        target = document.add_paragraph()
        target.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = target.add_run("Target role: " + tailored.target_title)
        run.bold = True
        run.font.size = Pt(10.5)

    _section_heading(document, "Professional Summary")
    document.add_paragraph(
        tailored.professional_summary or profile.professional_summary
    )

    skills = list(dict.fromkeys([*tailored.highlighted_skills, *profile.skills]))
    if skills:
        _section_heading(document, "Core Skills")
        document.add_paragraph(" • ".join(skills))

    if profile.experiences:
        _section_heading(document, "Professional Experience")
        for index, experience in enumerate(profile.experiences):
            header = document.add_paragraph()
            header.paragraph_format.keep_with_next = True
            run = header.add_run(experience.title)
            run.bold = True
            company_location = " | ".join(
                item for item in [experience.company, experience.location] if item
            )
            if company_location:
                header.add_run(f" — {company_location}")
            dates = " – ".join(
                item for item in [experience.start_date, experience.end_date] if item
            )
            if dates:
                header.add_run(f" | {dates}").italic = True
            for bullet in _experience_bullets(profile, tailored, index):
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.left_indent = Inches(0.18)
                paragraph.paragraph_format.first_line_indent = Inches(-0.14)
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.add_run(bullet)

    if profile.education:
        _section_heading(document, "Education")
        for education in profile.education:
            degree = " in ".join(
                item for item in [education.degree, education.field] if item
            )
            line = " | ".join(
                item
                for item in [
                    degree,
                    education.institution,
                    education.location,
                    " - ".join(
                        filter(None, [education.start_date, education.graduation_date])
                    ),
                ]
                if item
            )
            paragraph = document.add_paragraph(line)
            if paragraph.runs:
                paragraph.runs[0].bold = True
            for detail in education.details:
                document.add_paragraph(detail, style="List Bullet")

    if profile.certifications:
        _section_heading(document, "Certifications")
        document.add_paragraph(" • ".join(profile.certifications))
    if profile.projects:
        _section_heading(document, "Selected Projects")
        for project in profile.projects:
            document.add_paragraph(project, style="List Bullet")
    if profile.additional_information:
        _section_heading(document, "Additional Information")
        document.add_paragraph(" • ".join(profile.additional_information))
    document.save(destination)


def _pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "name": ParagraphStyle(
            "ResumeName",
            parent=base["Title"],
            fontName="Helvetica-Bold",
            fontSize=17,
            leading=19,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#17365D"),
            spaceAfter=2,
        ),
        "contact": ParagraphStyle(
            "Contact",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=8,
            leading=10,
            alignment=TA_CENTER,
            spaceAfter=4,
        ),
        "target": ParagraphStyle(
            "Target",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=10,
            leading=11,
            alignment=TA_CENTER,
            spaceAfter=4,
        ),
        "heading": ParagraphStyle(
            "Heading",
            parent=base["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=9.5,
            leading=11,
            textColor=colors.HexColor("#1F4E79"),
            borderWidth=0,
            borderPadding=0,
            spaceBefore=5,
            spaceAfter=2,
            keepWithNext=True,
        ),
        "body": ParagraphStyle(
            "Body",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=8.7,
            leading=10.4,
            spaceAfter=2,
        ),
        "job": ParagraphStyle(
            "Job",
            parent=base["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=9,
            leading=10.5,
            spaceBefore=2,
            spaceAfter=1,
            keepWithNext=True,
        ),
        "bullet": ParagraphStyle(
            "Bullet",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=8.5,
            leading=10.2,
            leftIndent=11,
            firstLineIndent=-7,
            bulletIndent=2,
            spaceAfter=1,
        ),
    }


def _section(story: list, styles: dict[str, ParagraphStyle], label: str) -> None:
    story.append(Paragraph(escape_paragraph(label.upper()), styles["heading"]))
    story.append(
        Paragraph(
            '<para backColor="#9EADBA" leading="1" spaceAfter="2"> </para>',
            ParagraphStyle("Rule", leading=1, spaceAfter=2),
        )
    )


def render_resume_pdf(
    profile: CandidateProfile,
    tailored: TailoredResume,
    destination: Path,
) -> None:
    render_ats_pdf(profile, tailored, destination)


def render_job_description_docx(job: JobListing, destination: Path) -> None:
    document = Document()
    _configure_docx(document)
    heading = document.add_paragraph(style="Title")
    heading.add_run(job.title)
    document.add_paragraph(job.company, style="Subtitle")
    metadata = " | ".join(
        item for item in [job.location, job.date_posted, job.source] if item
    )
    if metadata:
        document.add_paragraph(metadata)
    link = document.add_paragraph("Apply: ")
    _add_hyperlink(link, job.job_url, job.job_url)
    _section_heading(document, "Job Description")
    for block in job.description.split("\n"):
        text = block.strip().lstrip("•*- ").strip()
        if not text:
            continue
        if block.strip().startswith(("•", "-", "*")):
            document.add_paragraph(text, style="List Bullet")
        else:
            document.add_paragraph(text)
    document.save(destination)


def render_job_description_pdf(job: JobListing, destination: Path) -> None:
    styles = _pdf_styles()
    document = SimpleDocTemplate(
        str(destination),
        pagesize=letter,
        leftMargin=0.65 * inch,
        rightMargin=0.65 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
        title=f"{job.company} - {job.title}",
    )
    story: list = [
        Paragraph(escape_paragraph(job.title), styles["name"]),
        Paragraph(escape_paragraph(job.company), styles["target"]),
        Paragraph(
            escape_paragraph(
                " | ".join(
                    item for item in [job.location, job.date_posted, job.source] if item
                )
            ),
            styles["contact"],
        ),
        Paragraph(
            f'<link href="{html.escape(job.job_url, quote=True)}" color="#1F4E79">Apply to this role</link>',
            styles["body"],
        ),
        Spacer(1, 4),
    ]
    for block in job.description.split("\n"):
        stripped = block.strip()
        if not stripped:
            story.append(Spacer(1, 3))
            continue
        if stripped.startswith(("•", "-", "*")):
            story.append(
                Paragraph(
                    "• " + escape_paragraph(stripped.lstrip("•*- ")), styles["bullet"]
                )
            )
        else:
            story.append(Paragraph(escape_paragraph(stripped), styles["body"]))
    document.build(story)


def _write_apply_files(job: JobListing, folder: Path) -> Path:
    shortcut = folder / "APPLY_HERE.url"
    shortcut.write_text(f"[InternetShortcut]\nURL={job.job_url}\n", encoding="utf-8")
    html_file = folder / "APPLY_HERE.html"
    safe_url = html.escape(job.job_url, quote=True)
    safe_role = html.escape(f"{job.company} — {job.title}")
    html_file.write_text(
        "<!doctype html><html><head><meta charset='utf-8'><title>Apply</title></head>"
        "<body style='font-family:Arial,sans-serif;max-width:760px;margin:60px auto;padding:24px'>"
        f"<h1>{safe_role}</h1><p><a href='{safe_url}' target='_blank' rel='noopener'>"
        "Open the application page</a></p>"
        f"<p style='word-break:break-all'>{safe_url}</p></body></html>",
        encoding="utf-8",
    )
    (folder / "Application_Link.txt").write_text(job.job_url + "\n", encoding="utf-8")
    return shortcut


def _append_export_log(root: Path, job: JobListing, folder: Path, score: int) -> None:
    log_path = root / "application_packages.csv"
    exists = log_path.exists()
    with log_path.open("a", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        if not exists:
            writer.writerow(
                [
                    "created_at",
                    "company",
                    "role",
                    "location",
                    "date_posted",
                    "match_score",
                    "url",
                    "folder",
                ]
            )
        writer.writerow(
            [
                datetime.now(timezone.utc).isoformat(),
                job.company,
                job.title,
                job.location,
                job.date_posted,
                score,
                job.job_url,
                str(folder),
            ]
        )


def application_folder_name(profile: CandidateProfile, job: JobListing) -> str:
    # Keep Windows paths manageable; full names remain inside the documents.
    return "_".join(
        [
            slug_component(profile.contact.name, "Candidate", max_length=24),
            slug_component(job.company, "Company", max_length=36),
            slug_component(job.title, "Role", max_length=52),
        ]
    )


def create_application_package(
    output_root: str | Path,
    profile: CandidateProfile,
    job: JobListing,
    tailored: TailoredResume,
    validation: DraftValidation,
    selected_keywords: list[str],
    requirements: RequirementAnalysis | None = None,
    initial_draft: TailoredResume | None = None,
    destination_folder: Path | None = None,
) -> ExportResult:
    root = Path(output_root).expanduser().resolve()
    root.mkdir(parents=True, exist_ok=True)
    folder_name = application_folder_name(profile, job)
    if destination_folder is None:
        folder = ensure_unique_folder(root, folder_name)
        folder.mkdir(parents=False, exist_ok=False)
    else:
        folder = destination_folder.resolve()
        if folder.parent != root or not folder.is_dir():
            raise ValueError(
                "The application folder must be an existing direct child of the output directory."
            )

    changes = tailoring_changes(profile, tailored)
    from .evidence_tailoring import ALIGNED, READY, write_evidence_report

    report = tailored.tailoring_report or {
        "status": "Needs review: JD alignment not checked",
        "issues": ["No completed requirement-to-evidence review is available."],
    }
    write_evidence_report(folder, report)
    resume_stem = (
        "Tailored_Resume"
        if report["status"] == READY
        else "Resume_Reviewed"
        if report["status"] == ALIGNED
        else "Resume_Needs_Review"
    )
    resume_docx = folder / f"{resume_stem}.docx"
    resume_pdf = folder / f"{resume_stem}.pdf"
    jd_docx = folder / "Job_Description.docx"
    jd_pdf = folder / "Job_Description.pdf"
    render_resume_docx(profile, tailored, resume_docx)
    pdf_check = render_ats_pdf(profile, tailored, resume_pdf)
    write_json(folder / "ATS_Readability_Check.json", pdf_check)
    (folder / "Resume_Plain_Text.txt").write_text(
        "\n\n".join(
            ("- " if kind == "bullet" else "") + text
            for kind, text in resume_blocks(profile, tailored)
        ),
        encoding="utf-8",
    )
    render_job_description_docx(job, jd_docx)
    render_job_description_pdf(job, jd_pdf)
    apply_link = _write_apply_files(job, folder)
    write_json(folder / "Tailoring_Changes.json", changes)
    write_json(folder / "Source_Profile.json", profile.model_dump())
    write_json(folder / "Final_Tailoring.json", tailored.model_dump())
    write_json(
        folder / "Sponsorship_Report.json",
        {
            "role_status": job.sponsorship_status,
            "role_evidence": job.sponsorship_evidence,
            "source": job.sponsorship_source,
            "checked_at": job.sponsorship_checked_at,
            "company_history": job.sponsor_history,
            "notice": "Employer history is not a promise of sponsorship for this role. Confirm eligibility and terms with the employer.",
        },
    )
    if initial_draft is not None:
        write_json(folder / "Draft_Before_Audit.json", initial_draft.model_dump())
    if requirements is not None:
        write_json(folder / "JD_Requirements.json", requirements.model_dump())
    score = (
        tailored.match_score
        if tailored.match_rationale.strip() and requirements is not None
        else None
    )
    write_json(
        folder / "Match_Report.json",
        {
            "job": job.model_dump(),
            "match_score": score,
            "score_notice": "Model estimate, not an ATS score; omitted when no rationale or requirements analysis exists.",
            "match_rationale": tailored.match_rationale,
            "ats_keywords_used": tailored.ats_keywords_used,
            "missing_requirements": tailored.missing_requirements,
            "selected_search_keywords": selected_keywords,
            "validation": validation.model_dump(),
            "pdf_readability": pdf_check,
            "tailoring_changes": changes,
            "notice": "Review all generated material before applying. The source resume remains the authority.",
        },
    )
    _append_export_log(root, job, folder, score)
    return ExportResult(
        folder=str(folder),
        resume_docx=str(resume_docx),
        resume_pdf=str(resume_pdf),
        job_description_docx=str(jd_docx),
        job_description_pdf=str(jd_pdf),
        apply_link=str(apply_link),
        tailoring_status=changes["status"],
    )
